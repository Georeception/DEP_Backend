from django.contrib import admin
from django.contrib.auth.admin import UserAdmin
from django.utils.translation import gettext_lazy as _
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime
from .models import (
    User, News, NewsCategory, Event, EventCategory, EventRegistration,
    Gallery, GalleryCategory, NationalLeadership, LeadershipPosition,
    Donation, Product, ProductCategory, Order, OrderItem,
    MembershipPlan, Membership, NewsletterSubscription
)
from .models.locations import County, Constituency, Ward
from .models.shop import PickupLocation

# Location Admin
@admin.register(County)
class CountyAdmin(admin.ModelAdmin):
    list_display = ('name', 'code')
    search_fields = ('name', 'code')
    ordering = ('name',)

@admin.register(Constituency)
class ConstituencyAdmin(admin.ModelAdmin):
    list_display = ('name', 'code', 'county')
    list_filter = ('county',)
    search_fields = ('name', 'code', 'county__name')
    ordering = ('county', 'name')

@admin.register(Ward)
class WardAdmin(admin.ModelAdmin):
    list_display = ('name', 'code', 'constituency', 'get_county')
    list_filter = ('constituency__county', 'constituency')
    search_fields = ('name', 'code', 'constituency__name', 'constituency__county__name')
    ordering = ('constituency__county', 'constituency', 'name')

    def get_county(self, obj):
        return obj.constituency.county.name
    get_county.short_description = 'County'
    get_county.admin_order_field = 'constituency__county__name'

# User Admin
@admin.register(User)
class CustomUserAdmin(UserAdmin):
    list_display = ('email', 'first_name', 'last_name', 'phone_number', 'is_staff')
    list_filter = ('is_staff', 'is_superuser', 'is_active', 'groups')
    search_fields = ('email', 'first_name', 'last_name', 'phone_number')
    ordering = ('email',)
    
    fieldsets = (
        (None, {'fields': ('email', 'password')}),
        (_('Personal info'), {'fields': ('first_name', 'last_name', 'phone_number')}),
        (_('Permissions'), {
            'fields': ('is_active', 'is_staff', 'is_superuser', 'groups', 'user_permissions'),
        }),
        (_('Important dates'), {'fields': ('last_login', 'date_joined')}),
    )
    
    add_fieldsets = (
        (None, {
            'classes': ('wide',),
            'fields': ('email', 'password1', 'password2', 'first_name', 'last_name', 'phone_number'),
        }),
    )

# News Admin
@admin.register(NewsCategory)
class NewsCategoryAdmin(admin.ModelAdmin):
    list_display = ('name', 'slug',)
    prepopulated_fields = {'slug': ('name',)}
    search_fields = ('name',)

@admin.register(News)
class NewsAdmin(admin.ModelAdmin):
    list_display = ('title', 'category', 'author', 'is_published', 'created_at',)
    list_filter = ('is_published', 'category', 'created_at',)
    search_fields = ('title', 'content',)
    prepopulated_fields = {'slug': ('title',)}
    date_hierarchy = 'created_at'

# Event Admin
@admin.register(EventCategory)
class EventCategoryAdmin(admin.ModelAdmin):
    list_display = ('name', 'slug',)
    prepopulated_fields = {'slug': ('name',)}
    search_fields = ('name',)

@admin.register(Event)
class EventAdmin(admin.ModelAdmin):
    list_display = ('title', 'category', 'start_date', 'end_date', 'is_published',)
    list_filter = ('is_published', 'category', 'start_date',)
    search_fields = ('title', 'description',)
    prepopulated_fields = {'slug': ('title',)}
    date_hierarchy = 'start_date'

@admin.register(EventRegistration)
class EventRegistrationAdmin(admin.ModelAdmin):
    list_display = ('user', 'event', 'registration_date', 'status',)
    list_filter = ('status', 'registration_date',)
    search_fields = ('user__email', 'event__title',)
    date_hierarchy = 'registration_date'

# Gallery Admin
@admin.register(GalleryCategory)
class GalleryCategoryAdmin(admin.ModelAdmin):
    list_display = ('name', 'slug',)
    prepopulated_fields = {'slug': ('name',)}
    search_fields = ('name',)

@admin.register(Gallery)
class GalleryAdmin(admin.ModelAdmin):
    list_display = ('title', 'category', 'uploaded_by', 'created_at',)
    list_filter = ('category', 'created_at',)
    search_fields = ('title', 'description',)
    date_hierarchy = 'created_at'

# Leadership Admin
@admin.register(LeadershipPosition)
class LeadershipPositionAdmin(admin.ModelAdmin):
    list_display = ('title', 'slug', 'order',)
    prepopulated_fields = {'slug': ('title',)}
    search_fields = ('title',)
    ordering = ('order',)

@admin.register(NationalLeadership)
class NationalLeadershipAdmin(admin.ModelAdmin):
    list_display = ('user', 'position', 'start_date', 'end_date', 'is_active',)
    list_filter = ('is_active', 'position', 'start_date',)
    search_fields = ('user__email', 'position__title',)
    date_hierarchy = 'start_date'

# Donation Admin
@admin.register(Donation)
class DonationAdmin(admin.ModelAdmin):
    list_display = ('donor_name', 'amount', 'payment_method', 'status', 'created_at',)
    list_filter = ('status', 'payment_method', 'created_at',)
    search_fields = ('donor_name', 'donor_email', 'transaction_id',)
    date_hierarchy = 'created_at'

# Shop Admin
@admin.register(ProductCategory)
class ProductCategoryAdmin(admin.ModelAdmin):
    list_display = ('name', 'slug',)
    prepopulated_fields = {'slug': ('name',)}
    search_fields = ('name',)

@admin.register(Product)
class ProductAdmin(admin.ModelAdmin):
    list_display = ('name', 'category', 'price', 'stock', 'is_featured',)
    list_filter = ('is_featured', 'category',)
    search_fields = ('name', 'description',)
    prepopulated_fields = {'slug': ('name',)}

@admin.register(Order)
class OrderAdmin(admin.ModelAdmin):
    list_display = ('order_number', 'user', 'status', 'total_amount', 'created_at',)
    list_filter = ('status', 'created_at',)
    search_fields = ('order_number', 'user__email',)
    date_hierarchy = 'created_at'

@admin.register(OrderItem)
class OrderItemAdmin(admin.ModelAdmin):
    list_display = ('order', 'product', 'quantity', 'price',)
    list_filter = ('order__status',)
    search_fields = ('order__order_number', 'product__name',)

# Membership Admin
@admin.register(MembershipPlan)
class MembershipPlanAdmin(admin.ModelAdmin):
    list_display = ('title', 'type', 'amount', 'is_active',)
    list_filter = ('is_active', 'type',)
    search_fields = ('title', 'description',)

@admin.register(Membership)
class MembershipAdmin(admin.ModelAdmin):
    list_display = ('get_member_name', 'membership_type', 'payment_status', 'payment_method', 'amount', 'application_status', 'created_at')
    list_filter = ('membership_type', 'payment_status', 'payment_method', 'application_status', 'created_at')
    search_fields = ('user__email', 'email', 'first_name', 'last_name', 'phone', 'national_id_number')
    date_hierarchy = 'created_at'
    actions = ['export_to_excel']

    def get_member_name(self, obj):
        if obj.user:
            return f"{obj.user.get_full_name()} ({obj.user.email})"
        return f"{obj.first_name} {obj.last_name} ({obj.email})"
    get_member_name.short_description = 'Member'

    def export_to_excel(self, request, queryset):
        """
        Export selected membership records to Excel with all form fields.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Membership Applications"

        # Define header style
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Define column headers - matching the form sections
        headers = [
            # Application Info
            'ID', 'Date Applied', 'Application Status', 'Membership Reference',
            # Section A: Personal Details
            'Full Name', 'First Name', 'Last Name', 'National ID Number',
            'Date of Birth', 'Age', 'Gender', 'Email', 'Phone', 'WhatsApp Number',
            'County of Residence', 'Constituency of Residence', 'Ward of Residence',
            'Sub-County', 'Postal Address',
            # Section B: DEP-BUS Membership
            'DEP-BUS Membership Status', 'Membership Status Confirmed',
            # Section C: Position Applied For
            'Membership Type', 'Membership Category', 'Amount', 'Payment Status',
            'Payment Method', 'Transaction ID',
            # Section D: Electoral Area
            'Electoral County', 'Electoral Constituency', 'Electoral Ward',
            # Section E: Special Category
            'Special Category', 'Disability Nature',
            # Section F: Education, Profession & Experience
            'Education Level', 'Professional Qualifications',
            'Current Occupation', 'Current Employer', 'Professional Experience',
            # Section G: Leadership & Community Engagement
            'Previous Leadership Position', 'Previous Leadership Details',
            'Organizations Participated', 'Community Contribution',
            # Section H: Motivation for DEP-BUS Ticket
            'Motivation Reason', 'Key Issues', 'Devolution Understanding',
            'Contribution to Party',
            # Section I: Political & Electoral Information
            'Contested Before', 'Contested Details', 'Current Public Office',
            'Current Office Details', 'Bankruptcy Proceedings', 'Bankruptcy Details',
            'Criminal Proceedings', 'Criminal Proceedings Details',
            # Section J: Campaign & Organisational Capacity
            'Support Network', 'Volunteer Count', 'Campaign Structure',
            'Grassroots Strategy',
            # Section K: Commitment to DEP-BUS
            'Uphold Values', 'Peaceful Politics', 'Issue Based Politics',
            'Abide Nomination',
            # Section M: Document Attachments
            'ID Document', 'Passport Photo', 'Membership Proof',
            'Academic Certificates', 'Leadership Documents', 'Other Documents',
            # Admin Remarks
            'Membership Status Verified', 'Documents Verified', 'Admin Remarks',
        ]

        # Write headers
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Define data style
        data_alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

        # Write data rows
        for row_num, membership in enumerate(queryset, 2):
            # Application Info
            ws.cell(row=row_num, column=1, value=membership.id)
            ws.cell(row=row_num, column=2, value=membership.created_at.strftime('%Y-%m-%d %H:%M') if membership.created_at else '')
            ws.cell(row=row_num, column=3, value=membership.get_application_status_display() if membership.application_status else '')
            ws.cell(row=row_num, column=4, value=membership.membership_reference or '')

            # Section A: Personal Details
            ws.cell(row=row_num, column=5, value=membership.full_name or '')
            ws.cell(row=row_num, column=6, value=membership.first_name or '')
            ws.cell(row=row_num, column=7, value=membership.last_name or '')
            ws.cell(row=row_num, column=8, value=membership.national_id_number or '')
            ws.cell(row=row_num, column=9, value=membership.date_of_birth.strftime('%Y-%m-%d') if membership.date_of_birth else '')
            ws.cell(row=row_num, column=10, value=membership.age or '')
            ws.cell(row=row_num, column=11, value=membership.gender or '')
            ws.cell(row=row_num, column=12, value=membership.email or '')
            ws.cell(row=row_num, column=13, value=membership.phone or '')
            ws.cell(row=row_num, column=14, value=membership.whatsapp_number or '')
            ws.cell(row=row_num, column=15, value=membership.county.name if membership.county else '')
            ws.cell(row=row_num, column=16, value=membership.constituency.name if membership.constituency else '')
            ws.cell(row=row_num, column=17, value=membership.ward.name if membership.ward else '')
            ws.cell(row=row_num, column=18, value=membership.sub_county or '')
            ws.cell(row=row_num, column=19, value=membership.postal_address or '')

            # Section B: DEP-BUS Membership
            ws.cell(row=row_num, column=20, value=membership.get_dep_bus_membership_status_display() if membership.dep_bus_membership_status else '')
            ws.cell(row=row_num, column=21, value='Yes' if membership.membership_status_confirmed else 'No' if membership.membership_status_confirmed is not None else '')

            # Section C: Position Applied For
            ws.cell(row=row_num, column=22, value=membership.get_membership_type_display() if membership.membership_type else '')
            ws.cell(row=row_num, column=23, value=membership.get_membership_category_display() if membership.membership_category else '')
            ws.cell(row=row_num, column=24, value=str(membership.amount) if membership.amount else '')
            ws.cell(row=row_num, column=25, value=membership.get_payment_status_display() if membership.payment_status else '')
            ws.cell(row=row_num, column=26, value=membership.get_payment_method_display() if membership.payment_method else '')
            ws.cell(row=row_num, column=27, value=membership.transaction_id or '')

            # Section D: Electoral Area
            ws.cell(row=row_num, column=28, value=membership.electoral_county.name if membership.electoral_county else '')
            ws.cell(row=row_num, column=29, value=membership.electoral_constituency.name if membership.electoral_constituency else '')
            ws.cell(row=row_num, column=30, value=membership.electoral_ward.name if membership.electoral_ward else '')

            # Section E: Special Category
            ws.cell(row=row_num, column=31, value=membership.special_category or '')
            ws.cell(row=row_num, column=32, value=membership.disability_nature or '')

            # Section F: Education, Profession & Experience
            ws.cell(row=row_num, column=33, value=membership.education_level or '')
            ws.cell(row=row_num, column=34, value=membership.professional_qualifications or '')
            ws.cell(row=row_num, column=35, value=membership.current_occupation or '')
            ws.cell(row=row_num, column=36, value=membership.current_employer or '')
            ws.cell(row=row_num, column=37, value=membership.professional_experience or '')

            # Section G: Leadership & Community Engagement
            ws.cell(row=row_num, column=38, value=membership.get_previous_leadership_position_display() if membership.previous_leadership_position else '')
            ws.cell(row=row_num, column=39, value=membership.previous_leadership_details or '')
            ws.cell(row=row_num, column=40, value=membership.organizations_participated or '')
            ws.cell(row=row_num, column=41, value=membership.community_contribution or '')

            # Section H: Motivation for DEP-BUS Ticket
            ws.cell(row=row_num, column=42, value=membership.motivation_reason or '')
            ws.cell(row=row_num, column=43, value=membership.key_issues or '')
            ws.cell(row=row_num, column=44, value=membership.devolution_understanding or '')
            ws.cell(row=row_num, column=45, value=membership.contribution_to_party or '')

            # Section I: Political & Electoral Information
            ws.cell(row=row_num, column=46, value=membership.get_contested_before_display() if membership.contested_before else '')
            ws.cell(row=row_num, column=47, value=membership.contested_details or '')
            ws.cell(row=row_num, column=48, value=membership.get_current_public_office_display() if membership.current_public_office else '')
            ws.cell(row=row_num, column=49, value=membership.current_office_details or '')
            ws.cell(row=row_num, column=50, value=membership.get_bankruptcy_proceedings_display() if membership.bankruptcy_proceedings else '')
            ws.cell(row=row_num, column=51, value=membership.bankruptcy_details or '')
            ws.cell(row=row_num, column=52, value=membership.get_criminal_proceedings_display() if membership.criminal_proceedings else '')
            ws.cell(row=row_num, column=53, value=membership.criminal_proceedings_details or '')

            # Section J: Campaign & Organisational Capacity
            ws.cell(row=row_num, column=54, value=membership.support_network or '')
            ws.cell(row=row_num, column=55, value=membership.volunteer_count or '')
            ws.cell(row=row_num, column=56, value=membership.get_campaign_structure_display() if membership.campaign_structure else '')
            ws.cell(row=row_num, column=57, value=membership.grassroots_strategy or '')

            # Section K: Commitment to DEP-BUS
            ws.cell(row=row_num, column=58, value=membership.get_uphold_values_display() if membership.uphold_values else '')
            ws.cell(row=row_num, column=59, value=membership.get_peaceful_politics_display() if membership.peaceful_politics else '')
            ws.cell(row=row_num, column=60, value=membership.get_issue_based_politics_display() if membership.issue_based_politics else '')
            ws.cell(row=row_num, column=61, value=membership.get_abide_nomination_display() if membership.abide_nomination else '')

            # Section M: Document Attachments
            ws.cell(row=row_num, column=62, value=membership.id_document or '')
            ws.cell(row=row_num, column=63, value=membership.passport_photo or '')
            ws.cell(row=row_num, column=64, value=membership.membership_proof or '')
            ws.cell(row=row_num, column=65, value=membership.academic_certificates or '')
            ws.cell(row=row_num, column=66, value=membership.leadership_documents or '')
            ws.cell(row=row_num, column=67, value=membership.other_documents or '')

            # Admin Remarks
            ws.cell(row=row_num, column=68, value='Yes' if membership.membership_status_verified else 'No')
            ws.cell(row=row_num, column=69, value='Yes' if membership.documents_verified else 'No')
            ws.cell(row=row_num, column=70, value=membership.admin_remarks or '')

        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        # Generate response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename=membership_applications_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        wb.save(response)
        return response

    export_to_excel.short_description = "Export selected to Excel"

@admin.register(NewsletterSubscription)
class NewsletterSubscriptionAdmin(admin.ModelAdmin):
    list_display = ('email', 'subscription_date', 'status', 'is_verified', 'last_updated')
    list_filter = ('status', 'is_verified', 'subscription_date')
    search_fields = ('email',)
    date_hierarchy = 'subscription_date'
    readonly_fields = ('subscription_date', 'last_updated')
    fieldsets = (
        (None, {
            'fields': ('email', 'status', 'is_verified')
        }),
        ('Verification', {
            'fields': ('verification_token',)
        }),
        ('Timestamps', {
            'fields': ('subscription_date', 'last_updated'),
            'classes': ('collapse',)
        }),
    )

@admin.register(PickupLocation)
class PickupLocationAdmin(admin.ModelAdmin):
    list_display = ('name', 'city', 'state', 'phone', 'email', 'is_active')
    list_filter = ('is_active', 'city', 'state')
    search_fields = ('name', 'address', 'city', 'state', 'phone', 'email')
    ordering = ('city', 'name')
    readonly_fields = ('created_at', 'updated_at')
    fieldsets = (
        ('Basic Information', {
            'fields': ('name', 'is_active')
        }),
        ('Location Details', {
            'fields': ('address', 'city', 'state', 'zip_code')
        }),
        ('Contact Information', {
            'fields': ('phone', 'email')
        }),
        ('Timestamps', {
            'fields': ('created_at', 'updated_at'),
            'classes': ('collapse',)
        }),
    )
